"""Barcha analitika ma'lumotlarini bitta, tartibli Excel (.xlsx) faylga eksport qilish.

Har varaq: 1-qator sarlavha, 2-qator do'kon va davr, 4-qatordan jadval (formatlangan ustunlar, filtr, muzlatilgan sarlavha).
"""
import io
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.repositories import tq
from app.models import CatalogProduct, CustomerProfile, Order, Tenant
from app.services.analytics import metrics
from app.services.analytics.common import DateRange
from app.services.marketing.profiles import segment_label
from app.services.storefront import STATUS_LABELS

BRAND = "5B21B6"
HEAD_FILL = PatternFill("solid", fgColor="EDE9FE")
ZEBRA = PatternFill("solid", fgColor="FAFAFC")
THIN = Side(style="thin", color="E4E4E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ustun turi → Excel formati va tekislash
FORMATS = {
    "text": (None, "left"), "int": ("#,##0", "right"), "money": ("#,##0", "right"), "float": ("0.00", "right"),
    "pct": ('0.0"%"', "right"), "delta": ('+0.0"%";-0.0"%";0"%"', "right"), "date": ("dd.mm.yyyy", "center"),
    "datetime": ("dd.mm.yyyy hh:mm", "center"), "ms": ('#,##0" ms"', "right"),
}
WIDTH = {"text": 22, "int": 12, "money": 16, "float": 10, "pct": 11, "delta": 12, "date": 12, "datetime": 17, "ms": 11}

PAY_LABELS = {"pending": "Kutilmoqda", "paid": "To'langan", "failed": "Muvaffaqiyatsiz", "refunded": "Qaytarilgan",
              "cancelled": "Bekor"}
METHOD_LABELS = {"cash": "Naqd", "card": "Karta", "payme": "Payme", "click": "Click"}


def _cell_value(v: Any, kind: str) -> Any:
    if v is None or v == "":
        return None
    if kind in ("date", "datetime"):
        if isinstance(v, str):
            try:
                v = datetime.fromisoformat(v.replace("Z", ""))
            except ValueError:
                return v
        if isinstance(v, datetime):
            return v.replace(tzinfo=None) if kind == "datetime" else v.date()
        return v
    if kind == "money":
        return round(v / 100) if isinstance(v, (int, float)) else v
    if kind in ("pct", "delta") and isinstance(v, (int, float)):
        return round(float(v), 1)
    return v


class Book:
    def __init__(self, tenant: Tenant, r: DateRange):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.subtitle = f"{tenant.name}  ·  {r.start_local:%d.%m.%Y} — {r.end_local:%d.%m.%Y}  ·  valyuta: {tenant.setting('currency')}"

    def sheet(self, title: str, heading: str, columns: list[tuple[str, str]], rows: list[list[Any]],
              note: str | None = None, filters: bool = True):
        ws = self.wb.create_sheet(title[:31])
        ncol = len(columns)
        ws.cell(row=1, column=1, value=heading).font = Font(bold=True, size=14, color=BRAND)
        ws.cell(row=2, column=1, value=self.subtitle).font = Font(size=10, color="71717A")
        if note:
            ws.cell(row=2, column=max(3, ncol), value=note).font = Font(size=9, italic=True, color="A1A1AA")
        ws.row_dimensions[1].height = 22
        hdr = 4
        for i, (name, kind) in enumerate(columns, 1):
            c = ws.cell(row=hdr, column=i, value=name)
            c.font = Font(bold=True, color="3B0764")
            c.fill = HEAD_FILL
            c.border = BORDER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[hdr].height = 30
        if not rows:
            ws.cell(row=hdr + 1, column=1, value="Ma'lumot yo'q").font = Font(italic=True, color="A1A1AA")
        for ri, row in enumerate(rows, hdr + 1):
            zebra = (ri - hdr) % 2 == 0
            for ci, (_, kind) in enumerate(columns, 1):
                v = _cell_value(row[ci - 1] if ci - 1 < len(row) else None, kind)
                c = ws.cell(row=ri, column=ci, value=v)
                fmt, align = FORMATS[kind]
                if fmt:
                    c.number_format = fmt
                c.alignment = Alignment(horizontal=align, vertical="top", wrap_text=kind == "text")
                c.border = BORDER
                if zebra:
                    c.fill = ZEBRA
        for ci, (name, kind) in enumerate(columns, 1):
            longest = max([len(str(name)) // 1.4] + [len(str(r[ci - 1])) for r in rows[:300]
                                                     if ci - 1 < len(r) and r[ci - 1] is not None and kind == "text"])
            width = max(WIDTH[kind], min(50, int(longest) + 3)) if kind == "text" else WIDTH[kind]
            ws.column_dimensions[get_column_letter(ci)].width = width
        ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
        if filters and rows:
            ws.auto_filter.ref = f"A{hdr}:{get_column_letter(ncol)}{hdr + len(rows)}"
        ws.sheet_view.showGridLines = False
        return ws

    def bytes(self) -> bytes:
        buf = io.BytesIO()
        self.wb.save(buf)
        return buf.getvalue()


def _money_sum(items: list[dict], currency: str) -> int:
    return sum(x["amount_minor"] for x in items if x["currency"] == currency)


async def build_workbook(db: AsyncSession, tenant: Tenant, r: DateRange) -> bytes:
    book = Book(tenant, r)
    cur = tenant.setting("currency")

    # 1. Umumiy — joriy va oldingi davr yonma-yon
    ov = await metrics.overview(db, tenant, r)
    k, p, d = ov["kpis"], ov["previous"], ov["deltas"]
    kpi_rows = [
        ("AI qidiruvlar", k["searches"], p["searches"], d.get("searches"), "int"),
        ("Noyob qidiruvchilar", k["searchers"], p["searchers"], d.get("searchers"), "int"),
        ("Savatga qo'shilgan (noyob qidiruv)", k["add_to_cart"], p["add_to_cart"], d.get("add_to_cart"), "int"),
        ("Xarid (noyob qidiruv)", k["purchases"], p["purchases"], d.get("purchases"), "int"),
        ("Konversiya", k["conversion_pct"], p["conversion_pct"], d.get("conversion_pct"), "pct"),
        (f"Qidiruvdan kelgan daromad, {cur}", _money_sum(k["attributed_revenue"], cur), _money_sum(p["attributed_revenue"], cur),
         d.get("attributed_revenue"), "money"),
        (f"Yo'qotilgan daromad (tashlangan savat), {cur}", _money_sum(k["lost_revenue"], cur), _money_sum(p["lost_revenue"], cur),
         d.get("lost_revenue"), "money"),
        ("Qondirilmagan talab (qidiruvlar)", k["unmet_demand"], p["unmet_demand"], d.get("unmet_demand"), "int"),
        ("O'rtacha javob vaqti", k["avg_latency_ms"], p["avg_latency_ms"], None, "ms"),
        ("Keshdan berilgan javoblar", k["from_cache"], p["from_cache"], None, "int"),
    ]
    ws = book.sheet("Umumiy", "Umumiy ko'rsatkichlar",
                    [("Ko'rsatkich", "text"), ("Joriy davr", "int"), ("Oldingi davr", "int"), ("O'zgarish", "delta")],
                    [[n, a, b, c] for n, a, b, c, _ in kpi_rows], note=f"Eksport: {datetime.now():%d.%m.%Y %H:%M}", filters=False)
    for i, (*_, kind) in enumerate(kpi_rows, 5):      # qator bo'yicha format (pul, foiz, ms)
        for col in (2, 3):
            c = ws.cell(row=i, column=col)
            c.value = _cell_value(c.value, kind)
            c.number_format = FORMATS[kind][0]
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = ws.column_dimensions["C"].width = 18

    book.sheet("Kunlik dinamika", "Kunlik dinamika",
               [("Sana", "date"), ("Qidiruvlar", "int"), ("Qondirilmagan", "int"), ("Savatga", "int"), ("Xarid", "int")],
               [[date.fromisoformat(s["date"]), s["searches"], s["unmet"], s["add_to_cart"], s["purchases"]] for s in ov["series"]])

    fn = await metrics.funnel(db, tenant, r)
    book.sheet("Voronka", "Konversiya voronkasi (noyob qidiruvlar bo'yicha)",
               [("Bosqich", "text"), ("Soni", "int"), ("Jamidan", "pct"), ("Oldingi bosqichdan", "pct"), ("Yo'qotish", "int")],
               [[s["label"], s["count"], s["pct_of_total"], s["pct_of_prev"], s["drop"]] for s in fn["steps"]], filters=False)

    dm = await metrics.demand(db, tenant, r, limit=5000)
    book.sheet("Qondirilmagan talab", "Qidirilgan, lekin katalogda topilmagan mahsulotlar",
               [("So'rov", "text"), ("Kategoriya", "text"), ("Rang", "text"), ("Tur", "text"), ("Qidiruvlar", "int"),
                ("Noyob xaridorlar", "int"), ("Birinchi marta", "datetime"), ("Oxirgi marta", "datetime"), ("Tavsiya", "text")],
               [[i["sample_query"], i["category"], i["color"], i["type"], i["search_count"], i["unique_sessions"],
                 i["first_seen"], i["last_seen"], i["ai_suggestion"]] for i in dm["items"]])

    ab = await metrics.abandoned_carts(db, tenant, r, limit=10000)
    book.sheet("Tashlangan savat (mahsulot)", "Tashlab ketilgan savatlar — mahsulotlar bo'yicha",
               [("Mahsulot ID", "text"), ("Nomi", "text"), ("Dona", "int"), ("Savatlar", "int"), (f"Summa, {cur}", "money")],
               [[x["product_id"], x["name"], x["qty"], x["carts"], x["amount_minor"]] for x in ab["products"]])
    book.sheet("Tashlangan savatlar", "Tashlab ketilgan savatlar",
               [("Xaridor", "text"), ("Mahsulotlar", "text"), (f"Summa, {cur}", "money"), ("Oxirgi faollik", "datetime"),
                ("Rasm qidiruvidan", "text")],
               [[c["customer"], ", ".join(f"{i.get('name') or i.get('id')} ×{i.get('qty')}" for i in c["items"]),
                 c["total_minor"], c["last_activity_at"], "ha" if c["from_search"] else "yo'q"] for c in ab["carts"]])

    sl = await metrics.searches_log(db, tenant, r, 1, 20000)
    book.sheet("Qidiruvlar", "Rasm qidiruvlari jurnali",
               [("Vaqt (UTC)", "datetime"), ("Kategoriya", "text"), ("Rang", "text"), ("Tur", "text"), ("So'rov", "text"),
                ("AI ishonchi", "float"), ("Moslik", "float"), ("Natijalar", "int"), ("Topildimi", "text"),
                ("Javob vaqti", "ms"), ("Xaridor", "text")],
               [[s["created_at"], s["category"], s["color"], s["type"], s["query"], s["confidence"], s["best_score"],
                 s["results"], "yo'q" if s["unmet"] else "ha", s["latency_ms"], s["customer"]] for s in sl["items"]])

    cc = await metrics.colors_categories(db, tenant, r)
    rows = ([["Rang", x["name"], x["count"], x["unmet"]] for x in cc["colors"]]
            + [["Kategoriya", x["name"], x["count"], x["unmet"]] for x in cc["categories"]]
            + [["Mahsulot turi", x["name"], x["count"], x["unmet"]] for x in cc["types"]])
    book.sheet("Rang va kategoriya", "Eng ko'p qidirilgan ranglar, kategoriyalar va turlar",
               [("O'lchov", "text"), ("Qiymat", "text"), ("Qidiruvlar", "int"), ("Topilmagan", "int")], rows)
    book.sheet("Soatlik faollik", "Qidiruvlar soat bo'yicha (do'kon vaqt zonasida)",
               [("Soat", "text"), ("Qidiruvlar", "int")], [[f"{h['hour']:02d}:00", h["count"]] for h in cc["hourly"]], filters=False)

    profiles = (await db.execute(tq(CustomerProfile, tenant.id))).scalars().all()
    seg_count: dict[str, int] = {}
    for pr in profiles:
        for s in pr.segments or []:
            seg_count[s] = seg_count.get(s, 0) + 1
    book.sheet("Auditoriya segmentlari", "Auditoriya segmentlari",
               [("Segment", "text"), ("Profillar soni", "int")],
               [[segment_label(s), n] for s, n in sorted(seg_count.items(), key=lambda x: -x[1])])
    identified = [pr for pr in profiles if pr.customer_ref]
    book.sheet("Xaridorlar (customer_ref)", "Tanilgan xaridorlar va qiziqishlari",
               [("customer_ref", "text"), ("Kategoriyalar", "text"), ("Turlar", "text"), ("Ranglar", "text"),
                ("Qidiruv", "int"), ("Bosish", "int"), ("Savat", "int"), ("Xarid", "int"), (f"Daromad, {cur}", "money"),
                (f"Tashlangan, {cur}", "money"), ("Oxirgi faollik", "datetime")],
               [[pr.customer_ref, ", ".join(list((pr.interests or {}).get("category", {}))[:3]),
                 ", ".join(list((pr.interests or {}).get("type", {}))[:3]), ", ".join(list((pr.interests or {}).get("color", {}))[:3]),
                 pr.searches, pr.clicks, pr.carts, pr.purchases, pr.revenue_minor, pr.abandoned_minor, pr.last_seen]
                for pr in sorted(identified, key=lambda x: x.last_seen or datetime.min, reverse=True)[:20000]])

    if tenant.plan == "premium":
        orders = (await db.execute(tq(Order, tenant.id).where(Order.created_at >= r.start_utc, Order.created_at < r.end_utc)
                                   .order_by(Order.created_at.desc()))).scalars().all()
        book.sheet("Buyurtmalar", "Do'kon buyurtmalari",
                   [("Raqam", "text"), ("Sana (UTC)", "datetime"), ("Holat", "text"), ("To'lov", "text"), ("To'lov usuli", "text"),
                    ("Yetkazish", "text"), ("Xaridor", "text"), ("Telefon", "text"), ("Shahar", "text"),
                    (f"Mahsulotlar, {cur}", "money"), (f"Yetkazish, {cur}", "money"), (f"Jami, {cur}", "money"), ("Rasm qidiruvidan", "text")],
                   [[o.number, o.created_at, STATUS_LABELS.get(o.status, o.status), PAY_LABELS.get(o.payment_status, o.payment_status),
                     METHOD_LABELS.get(o.payment_method, o.payment_method), "Kuryer" if o.delivery_method == "courier" else "Olib ketish",
                     o.customer_name, o.phone, o.city, o.subtotal_minor, o.delivery_fee_minor, o.total_minor,
                     "ha" if o.search_id else "yo'q"] for o in orders])

    prods = (await db.execute(tq(CatalogProduct, tenant.id).where(CatalogProduct.is_deleted.is_(False))
                              .order_by(CatalogProduct.category, CatalogProduct.name))).scalars().all()
    book.sheet("Mahsulotlar", "Katalog",
               [("ID", "text"), ("Nomi", "text"), ("Kategoriya", "text"), ("Rang", "text"), ("Tur", "text"),
                (f"Narx, {cur}", "money"), (f"Eski narx, {cur}", "money"), ("Omborda", "int"), ("Sotuvda", "text"), ("Sotilgan", "int")],
               [[x.external_id, x.name, x.category, x.color, (x.attributes or {}).get("type"), x.price_minor, x.compare_at_minor,
                 x.stock_qty, "ha" if x.in_stock else "yo'q", x.sales_count] for x in prods])
    return book.bytes()
